"""
Generate the data files for the DS HS2026 exam AP02 (data ingestion, data
preparation & EDA) - a fictional Swiss online shop «PhoneMarkt».

Output (in the work package folder DS_HS2026_LNW_I_Examples/AP02):
  html/shop_page_1.html .. shop_page_3.html  product listing (3 pages, overlapping)
  html/product_detail.html                  detail page of one device (spec table)
  reviews.json                              saved response of a (fictional) reviews API
  smartphones_clean.csv                     cleaned listing data (formats unified,
                                            duplicates removed, outlier NOT removed)
  smartphone_data.xlsx                      sheets «Description» and «Data Sources»

Built-in data problems (used in the exam):
  - prices as "CHF 1'249.–" / "Preis auf Anfrage" (MNAR: price hidden above an
    internal threshold of CHF 1'900 -> the 8 most expensive devices)
  - storage as "256 GB" and "1 TB"
  - pagination overlap: the last 3 cards of page 1 and page 2 appear again at the
    top of the next page (exact duplicates)
  - one typo outlier: Samsung Galaxy S25 Ultra 512 GB «Titanium Silverblue» shown as
    CHF 12'990.– instead of CHF 1'299.– (same device in «Titanium Black» costs 1'299.–)
"""

import os, json, math, random
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

random.seed(42)
# Output goes to the AP02 work package folder (DS_HS2026_LNW_I_Examples/AP02)
BASE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                     '..', '..', 'DS_HS2026_LNW_I_Examples', 'AP02'))
HTML_DIR = os.path.join(BASE, 'html')
os.makedirs(HTML_DIR, exist_ok=True)

# ── catalogue ────────────────────────────────────────────────────────────────
# (brand, model, list price of the smallest storage, storages in GB,
#  battery mAh, screen inch, colours)
CATALOGUE = [
    ('Apple',    'iPhone 17 Pro Max',     1349, [256, 512, 1024], 4832, 6.9,  ['Cosmic Orange', 'Deep Blue']),
    ('Apple',    'iPhone 17 Pro',         1199, [256, 512, 1024], 4252, 6.3,  ['Silver', 'Deep Blue']),
    ('Apple',    'iPhone Air',            1049, [256, 512, 1024], 3149, 6.5,  ['Space Black', 'Sky Blue']),
    ('Apple',    'iPhone 17',              849, [256, 512],       3692, 6.3,  ['Lavender', 'Black']),
    ('Apple',    'iPhone 16',              749, [128, 256, 512],  3561, 6.1,  ['Black', 'Teal']),
    ('Apple',    'iPhone 16e',             599, [128, 256, 512],  4005, 6.1,  ['White', 'Black']),
    ('Samsung',  'Galaxy Z Fold7',        1899, [256, 512, 1024], 4400, 8.0,  ['Blue Shadow', 'Jetblack']),
    ('Samsung',  'Galaxy S25 Ultra',      1149, [256, 512, 1024], 5000, 6.9,  ['Titanium Black', 'Titanium Silverblue']),
    ('Samsung',  'Galaxy Z Flip7',        1099, [256, 512],       4300, 6.9,  ['Blue Shadow', 'Coralred']),
    ('Samsung',  'Galaxy S25+',            999, [256, 512],       4900, 6.7,  ['Navy', 'Icyblue']),
    ('Samsung',  'Galaxy S25',             799, [128, 256],       4000, 6.2,  ['Navy', 'Mint']),
    ('Samsung',  'Galaxy A56 5G',          429, [128, 256],       5000, 6.7,  ['Awesome Graphite', 'Awesome Olive']),
    ('Samsung',  'Galaxy A36 5G',          339, [128, 256],       5000, 6.7,  ['Awesome Black', 'Awesome Lavender']),
    ('Samsung',  'Galaxy A16 5G',          199, [128],            5000, 6.7,  ['Black', 'Light Green']),
    ('Google',   'Pixel 10 Pro Fold',     1799, [256, 512, 1024], 5015, 8.0,  ['Moonstone', 'Jade']),
    ('Google',   'Pixel 10 Pro XL',       1149, [256, 512, 1024], 5200, 6.8,  ['Obsidian', 'Porcelain']),
    ('Google',   'Pixel 10 Pro',           999, [128, 256, 512],  4870, 6.3,  ['Obsidian', 'Moonstone']),
    ('Google',   'Pixel 10',               799, [128, 256],       4970, 6.3,  ['Indigo', 'Lemongrass']),
    ('Google',   'Pixel 9a',               449, [128, 256],       5100, 6.3,  ['Obsidian', 'Iris']),
    ('Xiaomi',   'Xiaomi 15 Ultra',       1299, [512, 1024],      5410, 6.73, ['Black', 'Silver Chrome']),
    ('Xiaomi',   'Xiaomi 15T Pro',         749, [256, 512, 1024], 5500, 6.83, ['Black', 'Mocha Gold']),
    ('Xiaomi',   'Redmi Note 14 Pro 5G',   299, [256, 512],       5110, 6.67, ['Midnight Black', 'Ocean Blue']),
    ('Xiaomi',   'Redmi 15 5G',            179, [128, 256],       7000, 6.9,  ['Titan Gray', 'Ripple Green']),
    ('OnePlus',  'OnePlus 13',             899, [256, 512],       6000, 6.82, ['Black Eclipse', 'Arctic Dawn']),
    ('OnePlus',  'OnePlus Nord 5',         399, [256, 512],       5200, 6.83, ['Phantom Grey', 'Marble Sands']),
    ('Fairphone', 'Fairphone (Gen. 6)',    549, [256],            4415, 6.31, ['Cloud White', 'Forest Green']),
    ('Motorola', 'razr 60 ultra',         1099, [512],            4700, 7.0,  ['Rio Red', 'Scarab']),
    ('Motorola', 'edge 60',                379, [256],            5200, 6.67, ['Gibraltar Sea', 'Shamrock']),
    ('Motorola', 'moto g75 5G',            249, [256],            5000, 6.78, ['Charcoal Grey', 'Aqua Blue']),
    ('Nothing',  'Phone (3)',              799, [256, 512],       5150, 6.67, ['White', 'Black']),
    ('Nothing',  'Phone (3a)',             349, [128, 256],       5000, 6.77, ['White', 'Blue']),
    ('Sony',     'Xperia 1 VII',          1399, [256, 512],       5000, 6.5,  ['Moss Green', 'Slate Black']),
    ('Sony',     'Xperia 10 VII',          429, [128],            5000, 6.1,  ['Charcoal Black', 'Turquoise']),
]


def storage_step(price, from_gb):
    """Surcharge for the next storage level."""
    if price < 500:
        return {128: 50, 256: 80, 512: 120}[from_gb]
    return {128: 100, 256: 150, 512: 250}[from_gb]


def chf(value):
    """Swiss price format: CHF 1'249.–"""
    return "CHF {:,}.–".format(int(round(value))).replace(',', "'")


def storage_txt(gb):
    return '1&nbsp;TB' if gb == 1024 else f'{gb}&nbsp;GB'


# ── build products ───────────────────────────────────────────────────────────
products = []
for brand, model, base, storages, battery, screen, colours in CATALOGUE:
    price = base
    for i, gb in enumerate(storages):
        if i > 0:
            price += storage_step(base, storages[i - 1])
        for colour in colours:
            products.append(dict(brand=brand, model=model, storage_gb=gb, color=colour,
                                 screen_inch=screen, battery_mah=battery,
                                 list_price=float(price)))

random.shuffle(products)                       # shop sorts by «Beliebtheit»
for i, p in enumerate(products):
    p['product_id'] = f'PM-{10001 + i}'

# shop price: some items on sale
for p in products:
    p['old_price'] = None
    price = p['list_price']
    if random.random() < 0.18:
        p['old_price'] = price
        price = round(price * (1 - random.choice([0.05, 0.08, 0.10, 0.15])) / 10) * 10 - 1
    p['price_chf'] = float(price)
    p['availability'] = random.choices(['Sofort lieferbar', 'Lieferbar in 2–5 Tagen', 'Vorbestellung'],
                                       weights=[0.7, 0.2, 0.1])[0]

# «Preis auf Anfrage» - price hidden above an internal threshold of CHF 1'900 (MNAR):
# exactly the 8 most expensive devices (Galaxy Z Fold7 / Pixel 10 Pro Fold, 512 GB and 1 TB)
PRICE_THRESHOLD = 1900
on_request = [p for p in products if p['list_price'] > PRICE_THRESHOLD]
for p in on_request:
    p['price_chf'] = None
    p['old_price'] = None

# typo outlier: Galaxy S25 Ultra 512 GB Titanium Silverblue
s25u = [p for p in products if p['model'] == 'Galaxy S25 Ultra' and p['storage_gb'] == 512]
for p in s25u:
    p['price_chf'] = 1299.0
    p['old_price'] = None
outlier = next(p for p in s25u if p['color'] == 'Titanium Silverblue')
outlier['price_display'] = 12990.0

# ── listing pages (with overlap = duplicates) ───────────────────────────────
n = len(products)
cut1, cut2 = n // 3, 2 * n // 3
OVERLAP = 3
pages = [products[:cut1],
         products[cut1 - OVERLAP:cut2],
         products[cut2 - OVERLAP:]]

CSS = """
    body { font-family: Arial, Helvetica, sans-serif; margin: 0; background: #f5f5f7; color: #1d1d1f; }
    header { background: #0b3d91; color: #fff; padding: 12px 24px; display: flex; justify-content: space-between; }
    header a { color: #fff; margin-left: 16px; text-decoration: none; }
    .cookie-banner { background: #ffe9a8; padding: 8px 24px; font-size: 13px; }
    .container { display: flex; gap: 24px; padding: 24px; }
    aside { width: 200px; font-size: 14px; }
    .grid { display: grid; grid-template-columns: repeat(auto-fill, minmax(220px, 1fr)); gap: 16px; flex: 1; }
    .product-card { background: #fff; border-radius: 8px; padding: 16px; box-shadow: 0 1px 3px rgba(0,0,0,.1); }
    .product-title { font-size: 16px; margin: 0 0 8px 0; }
    .brand { font-weight: bold; }
    .specs { padding-left: 18px; font-size: 13px; color: #555; }
    .price { font-size: 18px; font-weight: bold; color: #c00; }
    .price-old { font-size: 13px; color: #888; margin-right: 6px; }
    .availability { font-size: 12px; color: #2a7a2a; }
    .pagination { padding: 0 24px 24px 24px; }
    .pagination a, .pagination span { margin-right: 8px; }
    footer { background: #222; color: #aaa; padding: 16px 24px; font-size: 12px; }
"""


def card_html(p):
    price = p.get('price_display', p['price_chf'])
    if price is None:
        price_box = '<span class="price on-request">Preis auf Anfrage</span>'
    else:
        old = f'<span class="price-old"><del>{chf(p["old_price"])}</del></span>' if p['old_price'] else ''
        price_box = f'{old}<span class="price">{chf(price)}</span>'
    battery = "{:,}".format(p['battery_mah']).replace(',', "'")
    return f"""
      <div class="product-card" data-product-id="{p['product_id']}">
        <a class="product-link" href="/p/{p['product_id']}">
          <h2 class="product-title"><span class="brand">{p['brand']}</span> <span class="model">{p['model']}</span></h2>
        </a>
        <ul class="specs">
          <li class="storage">{storage_txt(p['storage_gb'])}</li>
          <li class="color">{p['color']}</li>
          <li class="screen">{p['screen_inch']}&quot; Display</li>
          <li class="battery">{battery}&nbsp;mAh</li>
        </ul>
        <div class="price-box">{price_box}</div>
        <span class="availability">{p['availability']}</span>
      </div>"""


brands = sorted({p['brand'] for p in products})
for page_no, items in enumerate(pages, start=1):
    nav = []
    for k in range(1, 4):
        nav.append(f'<span class="current">{k}</span>' if k == page_no
                   else f'<a href="shop_page_{k}.html">{k}</a>')
    if page_no < 3:
        nav.append(f'<a class="next" href="shop_page_{page_no + 1}.html">Weiter &raquo;</a>')
    filters = '\n'.join(f'          <li class="filter-brand"><input type="checkbox"> {b}</li>' for b in brands)
    html = f"""<!DOCTYPE html>
<html lang="de-CH">
<head>
  <meta charset="utf-8">
  <title>Smartphones kaufen – Seite {page_no} | PhoneMarkt</title>
  <meta name="description" content="Smartphones aller Marken günstig online kaufen bei PhoneMarkt.">
  <style>{CSS}  </style>
</head>
<body>
  <header>
    <div class="logo">PhoneMarkt</div>
    <nav><a href="/">Home</a><a href="/smartphones">Smartphones</a><a href="/zubehoer">Zubehör</a><a href="/konto">Mein Konto</a></nav>
  </header>
  <div class="cookie-banner">Diese Website verwendet Cookies. <a href="/datenschutz">Mehr erfahren</a></div>
  <div class="breadcrumb" style="padding: 12px 24px; font-size: 13px;">Home &gt; Smartphones &gt; Alle Smartphones</div>
  <h1 style="padding: 0 24px;">Smartphones <small>({len(products)} Artikel, sortiert nach Beliebtheit)</small></h1>
  <div class="container">
    <aside>
      <h3>Filter</h3>
      <h4>Marke</h4>
      <ul class="filters">
{filters}
      </ul>
    </aside>
    <main class="grid">{''.join(card_html(p) for p in items)}
    </main>
  </div>
  <div class="pagination">Seite: {' '.join(nav)}</div>
  <footer>&copy; 2026 PhoneMarkt (fiktiver Online-Shop für Unterrichtszwecke) · Alle Preise in CHF inkl. MwSt.</footer>
</body>
</html>
"""
    with open(os.path.join(HTML_DIR, f'shop_page_{page_no}.html'), 'w', encoding='utf-8') as f:
        f.write(html)

# ── detail page ──────────────────────────────────────────────────────────────
detail = next(p for p in products if p['model'] == 'Pixel 10 Pro' and p['storage_gb'] == 256
              and p['color'] == 'Obsidian')
SPEC_GROUPS = [
    ('Allgemein', [('Marke', 'Google'), ('Modell', 'Pixel 10 Pro'), ('Betriebssystem', 'Android 16'),
                   ('Erscheinungsdatum', '28.08.2025'), ('Gewicht', '207 g')]),
    ('Display', [('Bildschirmdiagonale', '6.3"'), ('Auflösung', '1280 x 2856 Pixel'),
                 ('Bildwiederholrate', '1–120 Hz')]),
    ('Akku & Leistung', [('Akkukapazität', "4'870 mAh"), ('Prozessor', 'Google Tensor G5'),
                         ('Arbeitsspeicher (RAM)', '16 GB'), ('Speicherkapazität', '256 GB')]),
]
spec_rows = []
for group, rows in SPEC_GROUPS:
    spec_rows.append(f'        <tr class="spec-group"><th colspan="2">{group}</th></tr>')
    for k, v in rows:
        spec_rows.append(f'        <tr><td>{k}</td><td>{v}</td></tr>')
history = [('01.09.2025', 1149.0), ('01.11.2025', 1099.0), ('01.01.2026', 1049.0),
           ('01.04.2026', 1029.0), ('01.07.2026', detail['price_chf'])]
history_rows = '\n'.join(f'        <tr><td>{d}</td><td>{chf(v)}</td></tr>' for d, v in history)
detail_html = f"""<!DOCTYPE html>
<html lang="de-CH">
<head>
  <meta charset="utf-8">
  <title>Google Pixel 10 Pro 256 GB Obsidian | PhoneMarkt</title>
  <style>{CSS}
    table {{ border-collapse: collapse; margin-bottom: 24px; }}
    th, td {{ text-align: left; padding: 4px 12px; border-bottom: 1px solid #ddd; font-size: 14px; }}
    .spec-group th {{ background: #e8eef9; }}
  </style>
</head>
<body>
  <header>
    <div class="logo">PhoneMarkt</div>
    <nav><a href="/">Home</a><a href="/smartphones">Smartphones</a><a href="/zubehoer">Zubehör</a><a href="/konto">Mein Konto</a></nav>
  </header>
  <div class="breadcrumb" style="padding: 12px 24px; font-size: 13px;">Home &gt; Smartphones &gt; Google &gt; Pixel 10 Pro</div>
  <main style="padding: 0 24px;">
    <h1 class="product-name" data-product-id="{detail['product_id']}">Google Pixel 10 Pro 256 GB Obsidian</h1>
    <div class="price-box"><span class="price">{chf(detail['price_chf'])}</span></div>

    <h2>Preisverlauf</h2>
    <table class="price-history">
      <thead><tr><th>Datum</th><th>Preis</th></tr></thead>
      <tbody>
{history_rows}
      </tbody>
    </table>

    <h2>Technische Daten</h2>
    <table class="spec-table">
      <tbody>
{chr(10).join(spec_rows)}
      </tbody>
    </table>

    <h2>Lieferumfang</h2>
    <ul class="box-content">
      <li>Smartphone</li>
      <li>USB-C-auf-USB-C-Kabel (1 m)</li>
      <li>SIM-Tool</li>
    </ul>
  </main>
  <footer>&copy; 2026 PhoneMarkt (fiktiver Online-Shop für Unterrichtszwecke) · Alle Preise in CHF inkl. MwSt.</footer>
</body>
</html>
"""
with open(os.path.join(HTML_DIR, 'product_detail.html'), 'w', encoding='utf-8') as f:
    f.write(detail_html)

# ── reviews.json (saved API response) ───────────────────────────────────────
BRAND_EFFECT = {'Apple': 0.15, 'Samsung': 0.05, 'Google': 0.10, 'Xiaomi': -0.05, 'OnePlus': 0.10,
                'Fairphone': 0.20, 'Motorola': -0.15, 'Nothing': 0.05, 'Sony': -0.05}
NEW_MODELS = {'iPhone 17 Pro Max', 'iPhone 17 Pro', 'iPhone Air', 'iPhone 17', 'Galaxy Z Fold7',
              'Pixel 10 Pro Fold', 'Pixel 10 Pro XL', 'Pixel 10 Pro', 'Pixel 10', 'Xiaomi 15T Pro',
              'Phone (3)', 'Fairphone (Gen. 6)', 'Xperia 10 VII'}
results = []
for p in products:
    is_new = p['model'] in NEW_MODELS
    if random.random() < (0.30 if is_new else 0.08):
        continue                                    # no reviews yet
    lp = p['list_price']
    score = 3.55 + 0.55 * math.tanh((lp - 600) / 350) + BRAND_EFFECT[p['brand']] + random.gauss(0, 0.28)
    score = min(5.0, max(1.0, round(score, 1)))
    count = int(random.lognormvariate(3.2 if is_new else 4.4, 0.8)) + 1
    results.append({'product_id': p['product_id'], 'review_count': count, 'avg_score': score})
# discontinued devices still in the review system (not in the shop anymore)
for k, (pid, score, cnt) in enumerate([('PM-09871', 4.1, 412), ('PM-09874', 3.6, 158),
                                       ('PM-09902', 4.4, 233), ('PM-09917', 2.9, 61)]):
    results.append({'product_id': pid, 'review_count': cnt, 'avg_score': score})
results.sort(key=lambda r: r['product_id'])
# response body of GET /v1/reviews?category=smartphones (list of records, readable with pd.read_json)
with open(os.path.join(BASE, 'reviews.json'), 'w', encoding='utf-8') as f:
    json.dump(results, f, ensure_ascii=False, indent=2)

# ── clean csv (formats unified, duplicates removed, outlier NOT removed) ────
df_clean = pd.DataFrame([{
    'product_id':   p['product_id'],
    'brand':        p['brand'],
    'model':        p['model'],
    'storage_gb':   p['storage_gb'],
    'color':        p['color'],
    'screen_inch':  p['screen_inch'],
    'battery_mah':  p['battery_mah'],
    'price_chf':    p.get('price_display', p['price_chf']),
    'availability': p['availability'],
} for p in products])
df_clean.to_csv(os.path.join(BASE, 'smartphones_clean.csv'), index=False, encoding='utf-8')

# ── xlsx with description ───────────────────────────────────────────────────
DESCRIPTION = [
    ('Datei', 'Spalte', 'Beschreibung', 'Beispiel'),
    ('html/shop_page_1..3.html', '–', 'Gespeicherte Produktlisten-Seiten des (fiktiven) Online-Shops PhoneMarkt, 3 Seiten mit Paginierung', '–'),
    ('html/product_detail.html', '–', 'Gespeicherte Produkt-Detailseite eines Geräts mit Preisverlauf und technischen Daten', '–'),
    ('reviews.json', 'product_id', 'Produkt-ID (Schlüssel zu den Shop-Daten)', 'PM-10023'),
    ('reviews.json', 'review_count', 'Anzahl Kundenbewertungen', '87'),
    ('reviews.json', 'avg_score', 'Durchschnittliche Kundenbewertung (1–5 Sterne)', '4.3'),
    ('smartphones_clean.csv', 'product_id', 'Eindeutige Produkt-ID im Shop', 'PM-10023'),
    ('smartphones_clean.csv', 'brand', 'Marke', 'Samsung'),
    ('smartphones_clean.csv', 'model', 'Modellbezeichnung', 'Galaxy S25 Ultra'),
    ('smartphones_clean.csv', 'storage_gb', 'Speicherkapazität in GB (1 TB = 1024 GB)', '512'),
    ('smartphones_clean.csv', 'color', 'Farbe', 'Titanium Black'),
    ('smartphones_clean.csv', 'screen_inch', 'Bildschirmdiagonale in Zoll', '6.9'),
    ('smartphones_clean.csv', 'battery_mah', 'Akkukapazität in mAh', '5000'),
    ('smartphones_clean.csv', 'price_chf', 'Verkaufspreis in CHF (leer = «Preis auf Anfrage»)', '1299.0'),
    ('smartphones_clean.csv', 'availability', 'Lieferstatus', 'Sofort lieferbar'),
]
SOURCES = [
    ('Datei', 'Format', 'Art der Datenquelle', 'Hinweis'),
    ('html/shop_page_1..3.html', 'HTML', 'Web Scraping (Snapshot der Website)', 'Lokal gespeichert – kein Internetzugriff nötig'),
    ('html/product_detail.html', 'HTML', 'Web Scraping (Snapshot der Website)', 'Lokal gespeichert – kein Internetzugriff nötig'),
    ('reviews.json', 'JSON', 'REST-API (gespeicherte Antwort)', 'PhoneMarkt Reviews API (fiktiv): Antwort von GET /v1/reviews?category=smartphones, abgerufen am 15.09.2026'),
    ('smartphones_clean.csv', 'CSV', 'Aufbereiteter Datensatz', 'Formate vereinheitlicht, Duplikate entfernt; inhaltlich nicht geprüft'),
]
wb = Workbook()
for idx, (title, rows, widths) in enumerate([('Description', DESCRIPTION, [26, 14, 70, 20]),
                                             ('Data Sources', SOURCES, [26, 10, 38, 60])]):
    ws = wb.active if idx == 0 else wb.create_sheet()
    ws.title = title
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2E4057')
    for i, w in enumerate(widths):
        ws.column_dimensions['ABCD'[i]].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(os.path.join(BASE, 'smartphone_data.xlsx'))

print(f'Products: {len(products)}  | listing rows incl. duplicates: {sum(len(p) for p in pages)}')
print(f'Preis auf Anfrage: {len(on_request)}  | reviews: {len(results)}')
